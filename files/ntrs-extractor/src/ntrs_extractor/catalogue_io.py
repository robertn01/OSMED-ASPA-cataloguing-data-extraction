"""Read/write the catalogue as a formatted XLSX workbook.

Layout (one workbook per document):
    - Summary      : counts by kind / sub-catalogue / qc_flag
    - A_tabular    : sub-catalogue for tables         (schema.CATALOGUE_COLUMNS)
    - B_graphical  : sub-catalogue for figures/charts (schema.CATALOGUE_COLUMNS)
    - All          : every entry (convenience view)

The three data sheets share the same columns so a reviewer can copy rows between
them and downstream code can read any one of them uniformly.
"""
from __future__ import annotations

import os
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .schema import CATALOGUE_COLUMNS, CatalogueEntry

_HEADER_FILL = PatternFill("solid", start_color="1F3864")
_HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=10)
_BODY_FONT = Font(name="Arial", size=10)
_FLAG_FILL = {
    "green": PatternFill("solid", start_color="C6EFCE"),
    "amber": PatternFill("solid", start_color="FFEB9C"),
    "red": PatternFill("solid", start_color="FFC7CE"),
}
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _write_sheet(ws, rows: list[dict]):
    ws.append(CATALOGUE_COLUMNS)
    for ci, _ in enumerate(CATALOGUE_COLUMNS, start=1):
        c = ws.cell(row=1, column=ci)
        c.fill, c.font = _HEADER_FILL, _HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BORDER
    flag_col = CATALOGUE_COLUMNS.index("qc_flag") + 1
    for r in rows:
        ws.append([r.get(c, "") for c in CATALOGUE_COLUMNS])
        rr = ws.max_row
        for ci in range(1, len(CATALOGUE_COLUMNS) + 1):
            cell = ws.cell(row=rr, column=ci)
            cell.font = _BODY_FONT
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        flag = str(r.get("qc_flag", "")).lower()
        if flag in _FLAG_FILL:
            ws.cell(row=rr, column=flag_col).fill = _FLAG_FILL[flag]
    # widths
    widths = {
        "catalogue_id": 14, "printed_page_number": 8, "pdf_page_number": 8,
        "kind": 8, "item_number": 12, "original_source": 24, "description": 44,
        "approx_rows": 8, "approx_cols": 8, "sub_catalogue": 12,
        "figure_subtype": 12, "detection_method": 18, "detection_confidence": 10,
        "extractable": 9, "extraction_status": 14, "output_path": 26,
        "qc_flag": 8, "qc_notes": 28, "reviewer": 9,
    }
    for i, col in enumerate(CATALOGUE_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 12)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(CATALOGUE_COLUMNS))}{ws.max_row}"


def _write_summary(ws, rows: list[dict], doc_title: str):
    ws["A1"] = "NTRS Corpus Catalogue - Summary"
    ws["A1"].font = Font(bold=True, size=13, name="Arial")
    ws["A2"] = "Document:"
    ws["B2"] = doc_title
    ws["A2"].font = Font(bold=True, name="Arial")
    ws["B2"].font = _BODY_FONT

    def count(pred):
        return sum(1 for r in rows if pred(r))

    stats = [
        ("Total catalogued items", len(rows)),
        ("Tables (A_tabular)", count(lambda r: r["sub_catalogue"] == "A_tabular")),
        ("Figures (B_graphical)", count(lambda r: r["sub_catalogue"] == "B_graphical")),
        ("Extractable items", count(lambda r: r["extractable"] in (True, "True", "TRUE"))),
        ("Data-bearing figures", count(
            lambda r: r["kind"] == "figure"
            and r["extractable"] in (True, "True", "TRUE"))),
        ("QC green", count(lambda r: str(r["qc_flag"]).lower() == "green")),
        ("QC amber", count(lambda r: str(r["qc_flag"]).lower() == "amber")),
        ("QC red", count(lambda r: str(r["qc_flag"]).lower() == "red")),
    ]
    r0 = 4
    for i, (label, val) in enumerate(stats):
        ws.cell(row=r0 + i, column=1, value=label).font = Font(bold=True, name="Arial")
        ws.cell(row=r0 + i, column=2, value=val).font = _BODY_FONT
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 40


def write_catalogue(entries: Iterable[CatalogueEntry], out_path: str,
                    doc_title: str = "") -> str:
    rows = [e.as_row() if isinstance(e, CatalogueEntry) else e for e in entries]
    wb = Workbook()
    _write_summary(wb.active, rows, doc_title)
    wb.active.title = "Summary"
    _write_sheet(wb.create_sheet("A_tabular"),
                 [r for r in rows if r["sub_catalogue"] == "A_tabular"])
    _write_sheet(wb.create_sheet("B_graphical"),
                 [r for r in rows if r["sub_catalogue"] == "B_graphical"])
    _write_sheet(wb.create_sheet("All"), rows)
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    return out_path


def read_catalogue(path: str, sheet: str = "All") -> list[dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet]
    header = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        rows.append(dict(zip(header, row)))
    return rows
