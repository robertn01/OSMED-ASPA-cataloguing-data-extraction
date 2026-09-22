"""Quality-control helpers shared by all passes.

Each pass produces a QC record (a plain dict). `flag_from_score` maps a 0-1
score onto the green/amber/red traffic light used in the catalogue, and
`write_qc_report` dumps a per-document QC workbook so reviewers have one place to
triage everything that needs a human.
"""
from __future__ import annotations

import os
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

GREEN, AMBER, RED = "green", "amber", "red"

_FILL = {
    GREEN: PatternFill("solid", start_color="C6EFCE"),
    AMBER: PatternFill("solid", start_color="FFEB9C"),
    RED: PatternFill("solid", start_color="FFC7CE"),
}


def flag_from_score(score: float, green: float = 0.75, amber: float = 0.45) -> str:
    if score >= green:
        return GREEN
    if score >= amber:
        return AMBER
    return RED


def table_qc(result) -> dict:
    """QC record for a Pass-2 TableResult."""
    if result is None:
        return {"pass": "tables", "qc_score": 0.0, "qc_flag": RED,
                "reason": "no table extracted"}
    flag = flag_from_score(result.qc_score)
    reasons = []
    if result.empty_ratio > 0.35:
        reasons.append(f"empty_ratio={result.empty_ratio}")
    if result.ragged:
        reasons.append("ragged rows")
    if result.n_cols < 2:
        reasons.append("single column")
    return {
        "pass": "tables", "method": result.method, "qc_score": result.qc_score,
        "qc_flag": flag, "n_rows": result.n_rows, "n_cols": result.n_cols,
        "reason": "; ".join(reasons) or "ok",
    }


def figure_qc(result) -> dict:
    """QC record for a Pass-3 FigureResult."""
    if result is None:
        return {"pass": "figures", "qc_score": 0.0, "qc_flag": RED,
                "reason": "no figure digitised"}
    # uncalibrated results are capped at amber regardless of trace quality
    flag = flag_from_score(result.qc_score)
    if not result.calibrated and flag == GREEN:
        flag = AMBER
    reasons = []
    if not result.calibrated:
        reasons.append("uncalibrated (pixel space) - supply CalibrationSpec")
    if result.n_points < 8:
        reasons.append("few points")
    return {
        "pass": "figures", "qc_score": result.qc_score, "qc_flag": flag,
        "calibrated": result.calibrated, "n_points": result.n_points,
        "n_series": len(result.series), "reason": "; ".join(reasons) or "ok",
    }


def write_qc_report(records: Iterable[dict], out_path: str,
                    doc_title: str = "") -> str:
    records = list(records)
    wb = Workbook()
    ws = wb.active
    ws.title = "QC"
    ws["A1"] = f"QC report - {doc_title}"
    ws["A1"].font = Font(bold=True, size=13)
    if not records:
        os.makedirs(os.path.dirname(out_path), exist_ok=True)
        wb.save(out_path)
        return out_path
    cols = ["catalogue_id", "item_number", "pass", "method", "qc_score",
            "qc_flag", "n_rows", "n_cols", "n_points", "n_series",
            "calibrated", "reason", "output_path"]
    ws.append([])
    ws.append(cols)
    hdr = ws.max_row
    for i in range(1, len(cols) + 1):
        ws.cell(row=hdr, column=i).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=hdr, column=i).fill = PatternFill("solid", start_color="1F3864")
    flag_col = cols.index("qc_flag") + 1
    for rec in records:
        ws.append([rec.get(c, "") for c in cols])
        rr = ws.max_row
        f = str(rec.get("qc_flag", "")).lower()
        if f in _FILL:
            ws.cell(row=rr, column=flag_col).fill = _FILL[f]
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(c) + 2)
    ws.freeze_panes = f"A{hdr + 1}"
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    return out_path
